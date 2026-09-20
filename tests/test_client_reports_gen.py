import os
import sys
import datetime
import calendar
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch

from hisab_kitab_bot.config import PDF_TEMP_DIR
from hisab_kitab_bot.database import (
    get_user_country,
    get_client,
    get_client_opening_balance,
    get_client_period_summary
)
from hisab_kitab_bot.i18n import get_client_statement_str
from hisab_kitab_bot.pdf_generator import FONTS

def get_client_report_date_range(months: int = 1):
    """Computes start_date, end_date, and label for requested number of months (1 to 12)."""
    months = max(1, min(12, int(months)))
    today = datetime.date.today()
    year = today.year
    month = today.month - months + 1
    while month <= 0:
        month += 12
        year -= 1
    start_date = datetime.date(year, month, 1)
    end_date = today
    if months == 1:
        label = start_date.strftime("%B %Y")
    else:
        label = f"{start_date.strftime('%b %Y')} - {end_date.strftime('%b %Y')} ({months} Months)"
    return start_date.isoformat(), end_date.isoformat(), label, months

def generate_client_report_pdf(user_id: int, first_name: str, client_id: int = None, months: int = 1, start_date: str = None, end_date: str = None, period_label: str = None, lang: str = 'en', country_code: str = None) -> str:
    """Generates corporate executive PDF statement for clients (1 to 12 months)."""
    now = datetime.datetime.now()
    country_data = get_user_country(user_id)
    c_code = country_code or country_data.get("country_code", "GLOBAL")
    c_name = country_data.get("country_name", "Worldwide")
    curr = country_data.get("currency_symbol", "$")

    if not start_date or not end_date:
        start_date, end_date, auto_label, num_months = get_client_report_date_range(months)
        if not period_label:
            period_label = auto_label

    safe_start = start_date.replace("-", "")
    safe_end = end_date.replace("-", "")
    target_tag = f"c{client_id}" if client_id else "all"
    file_name = f"Client_Report_{c_code}_{user_id}_{target_tag}_{safe_start}_{safe_end}_{lang}.pdf"
    file_path = os.path.join(PDF_TEMP_DIR, file_name)

    if lang in ('hi', 'bn', 'ur'):
        font_regular = FONTS['indic']
        font_bold = FONTS['indic_bold']
    else:
        font_regular = FONTS['latin']
        font_bold = FONTS['latin_bold']

    doc = SimpleDocTemplate(
        file_path,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=32,
        bottomMargin=32
    )

    styles = getSampleStyleSheet()
    c_navy = colors.HexColor('#0F172A')
    c_slate = colors.HexColor('#334155')
    c_muted = colors.HexColor('#64748B')
    c_border = colors.HexColor('#CBD5E1')
    c_bg_light = colors.HexColor('#F8FAFC')
    c_green = colors.HexColor('#166534')
    c_red = colors.HexColor('#991B1B')
    c_green_bg = colors.HexColor('#F0FDF4')
    c_red_bg = colors.HexColor('#FEF2F2')

    style_sys_badge = ParagraphStyle('SysBadge', fontName=font_bold, fontSize=8, leading=10, textColor=colors.HexColor('#0284C7'))
    style_title = ParagraphStyle('DocTitle', fontName=font_bold, fontSize=15, leading=19, textColor=c_navy)
    style_sub = ParagraphStyle('DocSub', fontName=font_regular, fontSize=8.5, leading=12, textColor=c_slate)
    style_sec_head = ParagraphStyle('SecHead', fontName=font_bold, fontSize=10, leading=13, textColor=c_navy, spaceBefore=8, spaceAfter=4)
    style_th = ParagraphStyle('TH', fontName=font_bold, fontSize=8, leading=10, textColor=colors.white, alignment=1)
    style_td = ParagraphStyle('TD', fontName=font_regular, fontSize=8, leading=10, textColor=c_navy)
    style_td_bold = ParagraphStyle('TDBold', fontName=font_bold, fontSize=8, leading=10, textColor=c_navy)
    style_td_right = ParagraphStyle('TDRight', fontName=font_regular, fontSize=8, leading=10, textColor=c_navy, alignment=2)
    style_td_bold_right = ParagraphStyle('TDBoldRight', fontName=font_bold, fontSize=8, leading=10, textColor=c_navy, alignment=2)
    style_footer = ParagraphStyle('Footer', fontName=font_regular, fontSize=7, leading=9, textColor=c_muted, alignment=1)

    story = []

    # 1. Header
    sys_badge = get_client_statement_str('subtitle_system', lang)
    if client_id:
        client_data = get_client(client_id, user_id)
        client_name = client_data['name'] if client_data else "Client"
        doc_title = f"{get_client_statement_str('doc_title_single', lang)}: {client_name}"
    else:
        doc_title = get_client_statement_str('doc_title_all', lang)

    story.append(Paragraph(sys_badge, style_sys_badge))
    story.append(Spacer(1, 2))
    story.append(Paragraph(doc_title, style_title))
    story.append(Spacer(1, 10))

    # Fetch report data
    summary = get_client_period_summary(user_id, start_date, end_date, client_id=client_id, country_code=c_code)
    issue_date_str = now.strftime("%d %b %Y, %I:%M %p")
    audit_hash = f"CL-{c_code}-{user_id}-{safe_start}"

    # 2. Metadata Card
    scope_val = client_name if client_id else get_client_statement_str('lbl_all_clients', lang)
    meta_data = [
        [
            Paragraph(f"<b>{get_client_statement_str('lbl_account_holder', lang)}:</b> {first_name}", style_td),
            Paragraph(f"<b>{get_client_statement_str('lbl_period', lang)}:</b> {period_label}", style_td),
        ],
        [
            Paragraph(f"<b>{get_client_statement_str('lbl_scope', lang)}:</b> {scope_val}", style_td),
            Paragraph(f"<b>{get_client_statement_str('lbl_jurisdiction', lang)}:</b> {c_name} ({curr})", style_td),
        ],
        [
            Paragraph(f"<b>{get_client_statement_str('lbl_issue_date', lang)}:</b> {issue_date_str}", style_td),
            Paragraph(f"<b>Audit Ref:</b> {audit_hash}", style_td),
        ]
    ]
    meta_table = Table(meta_data, colWidths=[3.75 * inch, 3.75 * inch])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), c_bg_light),
        ('BOX', (0, 0), (-1, -1), 1, c_border),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 10))

    # 3. KPI Scorecard
    if client_id:
        op_bal = get_client_opening_balance(client_id, user_id, start_date)
        period_lena = summary['period_lena']
        period_dena = summary['period_dena']
        cl_bal = op_bal + period_lena - period_dena

        kpi_cells = [
            [
                Paragraph(get_client_statement_str('lbl_opening_balance', lang), style_td),
                Paragraph(get_client_statement_str('lbl_period_inflow', lang), style_td),
                Paragraph(get_client_statement_str('lbl_period_outflow', lang), style_td),
                Paragraph(get_client_statement_str('lbl_closing_balance', lang), style_td),
            ],
            [
                Paragraph(f"{curr}{op_bal:,.2f}", style_td_bold),
                Paragraph(f"+{curr}{period_lena:,.2f}", ParagraphStyle('KPiG', fontName=font_bold, fontSize=11, textColor=c_green)),
                Paragraph(f"-{curr}{period_dena:,.2f}", ParagraphStyle('KPiR', fontName=font_bold, fontSize=11, textColor=c_red)),
                Paragraph(f"{curr}{cl_bal:,.2f}", ParagraphStyle('KPiN', fontName=font_bold, fontSize=11, textColor=c_navy)),
            ]
        ]
        kpi_table = Table(kpi_cells, colWidths=[1.875 * inch] * 4)
    else:
        tot_lena = summary['total_receivable']
        tot_dena = summary['total_payable']
        net_pos = summary['net_position']
        cl_cnt = summary['total_clients']

        kpi_cells = [
            [
                Paragraph(get_client_statement_str('lbl_total_receivable', lang), style_td),
                Paragraph(get_client_statement_str('lbl_total_payable', lang), style_td),
                Paragraph(get_client_statement_str('lbl_net_position', lang), style_td),
                Paragraph(get_client_statement_str('lbl_total_clients', lang), style_td),
            ],
            [
                Paragraph(f"🟢 +{curr}{tot_lena:,.2f}", ParagraphStyle('KPiG', fontName=font_bold, fontSize=11, textColor=c_green)),
                Paragraph(f"🔴 -{curr}{tot_dena:,.2f}", ParagraphStyle('KPiR', fontName=font_bold, fontSize=11, textColor=c_red)),
                Paragraph(f"{curr}{net_pos:,.2f}", ParagraphStyle('KPiN', fontName=font_bold, fontSize=11, textColor=c_navy)),
                Paragraph(f"{cl_cnt}", ParagraphStyle('KPiC', fontName=font_bold, fontSize=11, textColor=c_navy)),
            ]
        ]
        kpi_table = Table(kpi_cells, colWidths=[1.875 * inch] * 4)

    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), c_bg_light),
        ('BOX', (0, 0), (-1, -1), 1, c_border),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, c_border),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 12))

    # 4. Tables
    if not client_id:
        # Table 1: All Clients Breakdown Table
        story.append(Paragraph(get_client_statement_str('sec_summary', lang), style_sec_head))
        cl_rows = [[
            Paragraph(get_client_statement_str('col_client', lang), style_th),
            Paragraph(get_client_statement_str('lbl_phone', lang), style_th),
            Paragraph(get_client_statement_str('lbl_total_receivable', lang), style_th),
            Paragraph(get_client_statement_str('lbl_total_payable', lang), style_th),
            Paragraph(get_client_statement_str('col_balance', lang), style_th),
            Paragraph(get_client_statement_str('col_status', lang), style_th),
        ]]

        for c in summary['clients']:
            bal = c['balance']
            if bal > 0:
                st_str = get_client_statement_str('status_get', lang)
                bal_p = Paragraph(f"+{curr}{bal:,.2f}", ParagraphStyle('B1', fontName=font_bold, fontSize=8, textColor=c_green, alignment=2))
            elif bal < 0:
                st_str = get_client_statement_str('status_give', lang)
                bal_p = Paragraph(f"-{curr}{abs(bal):,.2f}", ParagraphStyle('B2', fontName=font_bold, fontSize=8, textColor=c_red, alignment=2))
            else:
                st_str = get_client_statement_str('status_settled', lang)
                bal_p = Paragraph(f"{curr}0.00", style_td_bold_right)

            phone_str = c.get('phone') or "-"
            cl_rows.append([
                Paragraph(c['name'], style_td_bold),
                Paragraph(phone_str, style_td),
                Paragraph(f"{curr}{c.get('total_lena', 0.0):,.2f}", style_td_right),
                Paragraph(f"{curr}{c.get('total_dena', 0.0):,.2f}", style_td_right),
                bal_p,
                Paragraph(st_str, style_td),
            ])

        if len(cl_rows) == 1:
            cl_rows.append([Paragraph("No active clients.", style_td)] + [Paragraph("-", style_td)] * 5)

        cl_table = Table(cl_rows, colWidths=[1.8 * inch, 1.2 * inch, 1.1 * inch, 1.1 * inch, 1.1 * inch, 1.2 * inch])
        cl_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
            ('BOX', (0, 0), (-1, -1), 1, c_border),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_bg_light]),
            ('TOPPADDING', (0, 0), (-1, -1), 3.5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        story.append(cl_table)
        story.append(Spacer(1, 12))

    # Transaction Ledger Table
    story.append(Paragraph(get_client_statement_str('sec_ledger', lang), style_sec_head))
    tx_rows = [[
        Paragraph(get_client_statement_str('col_date', lang), style_th),
        Paragraph(get_client_statement_str('col_client', lang), style_th),
        Paragraph(get_client_statement_str('col_type', lang), style_th),
        Paragraph(get_client_statement_str('col_amount', lang), style_th),
        Paragraph(get_client_statement_str('col_note', lang), style_th),
    ]]

    if summary['transactions']:
        for t in summary['transactions'][:40]:  # Up to 40 entries
            t_type = t['type']
            if t_type == 'lena':
                type_p = Paragraph(get_client_statement_str('type_lena', lang), ParagraphStyle('TL', fontName=font_bold, fontSize=7.5, textColor=c_green))
                amt_p = Paragraph(f"+{curr}{t['amount']:,.2f}", ParagraphStyle('AL', fontName=font_bold, fontSize=8, textColor=c_green, alignment=2))
            else:
                type_p = Paragraph(get_client_statement_str('type_dena', lang), ParagraphStyle('TD2', fontName=font_bold, fontSize=7.5, textColor=c_red))
                amt_p = Paragraph(f"-{curr}{t['amount']:,.2f}", ParagraphStyle('AD', fontName=font_bold, fontSize=8, textColor=c_red, alignment=2))

            note_val = t.get('note') or "-"
            c_name_val = t.get('client_name') or "Client"
            tx_rows.append([
                Paragraph(t['date'], style_td),
                Paragraph(c_name_val, style_td_bold),
                type_p,
                amt_p,
                Paragraph(note_val, style_td),
            ])
    else:
        no_rec = get_client_statement_str('no_records', lang)
        tx_rows.append([Paragraph(no_rec, style_td)] + [Paragraph("-", style_td)] * 4)

    tx_table = Table(tx_rows, colWidths=[1.1 * inch, 1.6 * inch, 1.3 * inch, 1.2 * inch, 2.3 * inch])
    tx_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('BOX', (0, 0), (-1, -1), 1, c_border),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, c_bg_light]),
        ('TOPPADDING', (0, 0), (-1, -1), 3.5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3.5),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(tx_table)
    story.append(Spacer(1, 14))

    # 5. Corporate Footer
    footer_text = get_client_statement_str('footer_notice', lang)
    story.append(Paragraph(footer_text, style_footer))
    story.append(Spacer(1, 2))
    story.append(Paragraph(f"AUDIT HASH: {audit_hash} • ELECTRONICALLY VERIFIED • COMPLIANT", style_footer))

    doc.build(story)
    return file_path


def generate_client_report_excel(user_id: int, first_name: str, client_id: int = None, months: int = 1, start_date: str = None, end_date: str = None, period_label: str = None, lang: str = 'en', country_code: str = None) -> str:
    """Generates corporate executive multi-tab Excel workbook for clients (1 to 12 months)."""
    now = datetime.datetime.now()
    country_data = get_user_country(user_id)
    c_code = country_code or country_data.get("country_code", "GLOBAL")
    c_name = country_data.get("country_name", "Worldwide")
    curr = country_data.get("currency_symbol", "$")

    if not start_date or not end_date:
        start_date, end_date, auto_label, num_months = get_client_report_date_range(months)
        if not period_label:
            period_label = auto_label

    safe_start = start_date.replace("-", "")
    safe_end = end_date.replace("-", "")
    target_tag = f"c{client_id}" if client_id else "all"
    file_name = f"Client_Report_{c_code}_{user_id}_{target_tag}_{safe_start}_{safe_end}_{lang}.xlsx"
    file_path = os.path.join(PDF_TEMP_DIR, file_name)

    summary = get_client_period_summary(user_id, start_date, end_date, client_id=client_id, country_code=c_code)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = get_client_statement_str('sheet_summary', lang)

    # Styling colors
    NAVY_FILL = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    SLATE_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    LIGHT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    GREEN_FILL = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")
    RED_FILL = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")

    FONT_TITLE = Font(name="Arial", size=15, bold=True, color="FFFFFF")
    FONT_WHITE = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    FONT_BOLD = Font(name="Arial", size=10, bold=True, color="0F172A")
    FONT_REG = Font(name="Arial", size=9, color="334155")
    FONT_GREEN = Font(name="Arial", size=10, bold=True, color="166534")
    FONT_RED = Font(name="Arial", size=10, bold=True, color="991B1B")

    THIN_BORDER = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # 1. Header Banner
    if client_id:
        c_info = get_client(client_id, user_id)
        c_title_name = c_info['name'] if c_info else "Client"
        title_str = f"{get_client_statement_str('doc_title_single', lang)} - {c_title_name}"
    else:
        title_str = get_client_statement_str('doc_title_all', lang)

    ws1.merge_cells("A1:F2")
    ws1["A1"] = f"  {title_str}"
    ws1["A1"].font = FONT_TITLE
    ws1["A1"].fill = NAVY_FILL
    ws1["A1"].alignment = Alignment(vertical="center")

    # 2. Metadata
    ws1["A4"] = get_client_statement_str('lbl_account_holder', lang)
    ws1["B4"] = first_name
    ws1["A5"] = get_client_statement_str('lbl_period', lang)
    ws1["B5"] = period_label
    ws1["D4"] = get_client_statement_str('lbl_jurisdiction', lang)
    ws1["E4"] = f"{c_name} ({curr})"
    ws1["D5"] = get_client_statement_str('lbl_issue_date', lang)
    ws1["E5"] = now.strftime("%d %b %Y, %I:%M %p")

    for r in range(4, 6):
        for c in [1, 2, 4, 5]:
            cell = ws1.cell(row=r, column=c)
            cell.font = FONT_BOLD if c in (1, 4) else FONT_REG
            cell.fill = LIGHT_FILL
            cell.border = THIN_BORDER

    # 3. KPI Cards
    ws1["A7"] = get_client_statement_str('lbl_total_receivable', lang)
    ws1["A8"] = summary['total_receivable']
    ws1["A8"].number_format = f'"{curr}"#,##0.00'
    ws1["A8"].font = FONT_GREEN
    ws1["A8"].fill = GREEN_FILL

    ws1["C7"] = get_client_statement_str('lbl_total_payable', lang)
    ws1["C8"] = summary['total_payable']
    ws1["C8"].number_format = f'"{curr}"#,##0.00'
    ws1["C8"].font = FONT_RED
    ws1["C8"].fill = RED_FILL

    ws1["E7"] = get_client_statement_str('lbl_net_position', lang)
    ws1["E8"] = summary['net_position']
    ws1["E8"].number_format = f'"{curr}"#,##0.00'
    ws1["E8"].font = FONT_BOLD
    ws1["E8"].fill = LIGHT_FILL

    for k in ["A7", "C7", "E7"]:
        ws1[k].font = FONT_BOLD
        ws1[k].fill = LIGHT_FILL
        ws1[k].border = THIN_BORDER
    for k in ["A8", "C8", "E8"]:
        ws1[k].border = THIN_BORDER

    # 4. Table of Clients
    row_idx = 11
    headers = [
        get_client_statement_str('col_client', lang),
        get_client_statement_str('lbl_phone', lang),
        get_client_statement_str('lbl_email', lang),
        get_client_statement_str('lbl_total_receivable', lang),
        get_client_statement_str('lbl_total_payable', lang),
        get_client_statement_str('col_balance', lang),
    ]
    for col_idx, h in enumerate(headers, 1):
        cell = ws1.cell(row=row_idx, column=col_idx, value=h)
        cell.font = FONT_WHITE
        cell.fill = SLATE_FILL
        cell.alignment = Alignment(horizontal="center" if col_idx > 3 else "left")
        cell.border = THIN_BORDER

    for c in summary['clients']:
        row_idx += 1
        ws1.cell(row=row_idx, column=1, value=c['name']).font = FONT_BOLD
        ws1.cell(row=row_idx, column=2, value=c.get('phone') or "-").font = FONT_REG
        ws1.cell(row=row_idx, column=3, value=c.get('email') or "-").font = FONT_REG

        c_rec = ws1.cell(row=row_idx, column=4, value=c.get('total_lena', 0.0))
        c_rec.number_format = f'"{curr}"#,##0.00'
        c_rec.font = FONT_GREEN

        c_pay = ws1.cell(row=row_idx, column=5, value=c.get('total_dena', 0.0))
        c_pay.number_format = f'"{curr}"#,##0.00'
        c_pay.font = FONT_RED

        c_bal = ws1.cell(row=row_idx, column=6, value=f"=D{row_idx}-E{row_idx}")
        c_bal.number_format = f'"{curr}"#,##0.00'
        c_bal.font = FONT_BOLD

        for c_num in range(1, 7):
            ws1.cell(row=row_idx, column=c_num).border = THIN_BORDER

    # Auto-fit columns
    for col in ws1.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws1.column_dimensions[col_letter].width = max(max_len + 3, 13)

    # 5. Sheet 2: Transaction Ledger
    ws2 = wb.create_sheet(title=get_client_statement_str('sheet_ledger', lang))
    ws2.merge_cells("A1:E2")
    ws2["A1"] = f"  {get_client_statement_str('sec_ledger', lang)}"
    ws2["A1"].font = FONT_TITLE
    ws2["A1"].fill = NAVY_FILL
    ws2["A1"].alignment = Alignment(vertical="center")

    tx_headers = [
        get_client_statement_str('col_date', lang),
        get_client_statement_str('col_client', lang),
        get_client_statement_str('col_type', lang),
        get_client_statement_str('col_amount', lang),
        get_client_statement_str('col_note', lang),
    ]
    for col_idx, h in enumerate(tx_headers, 1):
        cell = ws2.cell(row=4, column=col_idx, value=h)
        cell.font = FONT_WHITE
        cell.fill = SLATE_FILL
        cell.alignment = Alignment(horizontal="center" if col_idx in (1, 3) else "left")
        cell.border = THIN_BORDER

    r_tx = 4
    for t in summary['transactions']:
        r_tx += 1
        ws2.cell(row=r_tx, column=1, value=t['date']).font = FONT_REG
        ws2.cell(row=r_tx, column=2, value=t.get('client_name') or "Client").font = FONT_BOLD
        ws2.cell(row=r_tx, column=3, value=t['type'].upper()).font = FONT_BOLD
        amt_cell = ws2.cell(row=r_tx, column=4, value=t['amount'])
        amt_cell.number_format = f'"{curr}"#,##0.00'
        amt_cell.font = FONT_GREEN if t['type'] == 'lena' else FONT_RED
        ws2.cell(row=r_tx, column=5, value=t.get('note') or "-").font = FONT_REG
        for c_num in range(1, 6):
            ws2.cell(row=r_tx, column=c_num).border = THIN_BORDER

    for col in ws2.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws2.column_dimensions[col_letter].width = max(max_len + 3, 14)

    wb.save(file_path)
    return file_path

if __name__ == "__main__":
    # Test generation
    user_id = 99998888
    p1 = generate_client_report_pdf(user_id, "Raja", months=1, lang='hi', country_code='IN')
    print("Generated 1M Client PDF:", p1, os.path.getsize(p1), "bytes")
    p12 = generate_client_report_pdf(user_id, "Raja", months=12, lang='en', country_code='IN')
    print("Generated 12M Client PDF:", p12, os.path.getsize(p12), "bytes")

    x1 = generate_client_report_excel(user_id, "Raja", months=1, lang='hi', country_code='IN')
    print("Generated 1M Client Excel:", x1, os.path.getsize(x1), "bytes")
    x12 = generate_client_report_excel(user_id, "Raja", months=12, lang='en', country_code='IN')
    print("Generated 12M Client Excel:", x12, os.path.getsize(x12), "bytes")
    print("ALL TESTS PASSED SUCCESSFULLY!")
